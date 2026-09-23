from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
MASTER = ROOT / "portfolio" / "resume" / "Palla_Janakiram_Data_Analyst_Resume.docx"
OUT = ROOT / "portfolio" / "job_search" / "resumes"
OUT.mkdir(parents=True, exist_ok=True)


def replace_paragraph(paragraph, text, size=8.55):
    for run in paragraph.runs:
        run._element.getparent().remove(run._element)
    run = paragraph.add_run(text)
    run.font.size = Pt(size)


def replace_skills(paragraph, groups):
    for run in paragraph.runs:
        run._element.getparent().remove(run._element)
    for index, (label, values) in enumerate(groups):
        lead = paragraph.add_run(f"{label}: ")
        lead.bold = True
        lead.font.size = Pt(8.35)
        body = paragraph.add_run(values + ("  |  " if index < len(groups) - 1 else ""))
        body.font.size = Pt(8.35)


CONFIGS = {
    "Target_Data_Analyst": {
        "summary": "Data Analyst with 2+ years of experience extracting, cleaning, transforming, validating and analyzing structured operational data. Uses SQL, Python/Pandas, Excel, PostgreSQL and Power BI to develop KPIs, dimensional models, dashboards and clear business insights. Portfolio work adds Spark/PySpark, Airflow, Git and validated Generative AI-assisted analytical workflows.",
        "skills": [
            ("Data Analytics", "SQL, Python, Pandas, NumPy, Excel, data cleaning, transformation, validation, EDA, statistical analysis, KPI analysis, business analysis"),
            ("Business Intelligence", "Power BI, DAX, Power Query, data visualization, dashboard development, dimensional modeling, data storytelling"),
            ("Data Platforms", "PostgreSQL, PySpark, Apache Spark, Delta Lake, dbt, medallion architecture, Airflow, incremental processing"),
            ("Tools", "Git, GitHub, ChatGPT, Microsoft Copilot; independently validated outputs"),
        ],
        "bullets": {
            9: "Analyze and validate high-volume structured operational datasets against specifications and business rules, identifying patterns, inconsistencies and data-quality risks.",
            10: "Execute consistency checks across complex attributes to detect exceptions, incomplete records and conflicting values, supporting accurate analytical outputs.",
            11: "Investigate root causes, document findings and coordinate correction workflows with cross-functional teams.",
            16: "Unified customer, transaction, deposit, lending, fraud, churn and branch data to answer business questions about behavior, trends, risk and operating performance.",
            17: "Cleaned, joined and validated 257,969 records; used SQL and Python/Pandas to calculate and reconcile KPIs across 5,000 customers and 120,000 transactions.",
            18: "Built a six-page Power BI solution with dimensional modeling, DAX, synchronized filters and decision-focused data storytelling.",
            19: "Implemented a local PySpark/Delta Lake pipeline with incremental processing, dbt tests, automated data-quality checks, Git version control and an eight-task Airflow DAG design.",
        },
    },
    "Enphase_Analyst_II_Data_Analysis": {
        "summary": "Data Analyst with 2+ years of experience preparing, validating and analyzing structured operational datasets for reporting and decision support. Uses SQL, Python/Pandas, Excel, Power BI and PostgreSQL to build KPI dashboards, investigate discrepancies and communicate actionable insights. Portfolio work includes PySpark, automated data quality and independently validated analytical outputs.",
        "skills": [
            ("Analytics", "SQL, Python, Pandas, NumPy, Advanced Excel, data cleaning, transformation, validation, EDA, KPI analysis, root-cause analysis"),
            ("Reporting and BI", "Power BI, DAX, Power Query, dashboards, data visualization, dimensional modeling, data storytelling"),
            ("Data Platforms", "PostgreSQL, PySpark, Apache Spark, Delta Lake, dbt, incremental processing, automated data quality"),
            ("Tools", "Git, GitHub, ChatGPT, Microsoft Copilot; independently validated outputs"),
        ],
        "bullets": {
            9: "Analyze and validate high-volume structured operational datasets against specifications and business rules, maintaining accuracy across recurring quality workflows.",
            10: "Monitor complex data attributes and perform consistency checks to identify exceptions, incomplete records and conflicting values.",
            11: "Investigate discrepancies and root causes, document findings and coordinate corrections with cross-functional teams.",
            12: "Extracted, cleaned, transformed and validated multi-source datasets using Python, Pandas and Excel, producing analysis-ready outputs and structured reports.",
            16: "Created banking KPI reporting across customers, transactions, deposits, lending, fraud, collections, churn and branches for operational and executive decision support.",
            17: "Used SQL and Python/Pandas to clean and reconcile 257,969 records, including 5,000 customers and 120,000 transactions.",
            18: "Built a six-page Power BI dashboard with DAX, Power Query, drill-down analysis and KPIs for portfolio risk, digital engagement and branch performance.",
            19: "Added a locally validated PySpark/Delta Lake pipeline with incremental processing, dbt tests, automated quality controls and Airflow DAG design.",
        },
    },
}


for filename, config in CONFIGS.items():
    doc = Document(MASTER)
    replace_paragraph(doc.paragraphs[5], config["summary"], 8.6)
    replace_skills(doc.paragraphs[7], config["skills"])
    for paragraph_index, text in config["bullets"].items():
        replace_paragraph(doc.paragraphs[paragraph_index], text)
    output = OUT / f"Palla_Janakiram_{filename}.docx"
    doc.save(output)
    print(output)
