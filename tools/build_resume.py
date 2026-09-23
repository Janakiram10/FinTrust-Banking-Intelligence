from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/"portfolio"/"resume"; OUT.mkdir(parents=True,exist_ok=True)
DOCX=OUT/"Palla_Janakiram_Data_Analyst_Resume.docx"

def hyperlink(p,text,url):
    part=p.part; rel=part.relate_to(url,"http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",is_external=True)
    h=OxmlElement("w:hyperlink"); h.set(qn("r:id"),rel); r=OxmlElement("w:r"); pr=OxmlElement("w:rPr")
    color=OxmlElement("w:color"); color.set(qn("w:val"),"0B5E75"); pr.append(color)
    u=OxmlElement("w:u"); u.set(qn("w:val"),"single"); pr.append(u); r.append(pr)
    t=OxmlElement("w:t"); t.text=text; r.append(t); h.append(r); p._p.append(h)

def heading(doc,text):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(2)
    r=p.add_run(text.upper()); r.bold=True; r.font.size=Pt(9.5); r.font.color.rgb=RGBColor(0x0B,0x3A,0x53)
    b=OxmlElement("w:pBdr"); bottom=OxmlElement("w:bottom"); bottom.set(qn("w:val"),"single"); bottom.set(qn("w:sz"),"4"); bottom.set(qn("w:color"),"8DB4C2"); b.append(bottom); p._p.get_or_add_pPr().append(b)

def bullet(doc,text):
    p=doc.add_paragraph(style="List Bullet"); p.paragraph_format.left_indent=Inches(.16); p.paragraph_format.first_line_indent=Inches(-.12)
    p.paragraph_format.space_after=Pt(.7); p.paragraph_format.line_spacing=1.0
    r=p.add_run(text); r.font.size=Pt(8.55); return p

def role(doc,title,org,dates):
    t=doc.add_table(rows=1,cols=2); t.autofit=False; t.columns[0].width=Inches(5.5); t.columns[1].width=Inches(1.45)
    t.cell(0,0).width=Inches(5.5); t.cell(0,1).width=Inches(1.45)
    for c in t.row_cells(0): c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p=t.cell(0,0).paragraphs[0]; p.paragraph_format.space_after=Pt(0); r=p.add_run(title); r.bold=True; r.font.size=Pt(9); p.add_run(f" | {org}").font.size=Pt(9)
    p=t.cell(0,1).paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.RIGHT; p.paragraph_format.space_after=Pt(0); r=p.add_run(dates); r.bold=True; r.font.size=Pt(8.6)
    for cell in t.row_cells(0):
        tcPr=cell._tc.get_or_add_tcPr(); mar=OxmlElement("w:tcMar")
        for side in ("top","left","bottom","right"):
            x=OxmlElement(f"w:{side}"); x.set(qn("w:w"),"0"); x.set(qn("w:type"),"dxa"); mar.append(x)
        tcPr.append(mar)

doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.42); sec.bottom_margin=Inches(.42); sec.left_margin=Inches(.55); sec.right_margin=Inches(.55)
styles=doc.styles; styles["Normal"].font.name="Arial"; styles["Normal"].font.size=Pt(8.4)
styles["Title"].font.name="Arial"; styles["Title"].font.color.rgb=RGBColor(0,0,0)

p=doc.add_paragraph(style="Title"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(0)
r=p.add_run("PALLA JANAKIRAM"); r.bold=True; r.font.size=Pt(18); r.font.color.rgb=RGBColor(0,0,0)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(1)
r=p.add_run("DATA ANALYST"); r.bold=True; r.font.size=Pt(10); r.font.color.rgb=RGBColor(0x0B,0x5E,0x75)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(1)
r=p.add_run("SQL  |  PYTHON  |  EXCEL  |  POWER BI"); r.bold=True; r.font.size=Pt(8.5); r.font.color.rgb=RGBColor(0x0B,0x3A,0x53)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(2)
p.add_run("Hyderabad, Telangana  |  +91 9866804339  |  janakiram.palla01@gmail.com  |  ").font.size=Pt(8)
hyperlink(p,"GitHub","https://github.com/Janakiram10")

heading(doc,"Professional Summary")
p=doc.add_paragraph("Data Analyst with 2+ years of experience extracting, cleaning, transforming, validating and analyzing structured operational datasets. Uses SQL, Python/Pandas, Excel, PostgreSQL, Power BI and DAX to develop KPIs, dimensional models, dashboards and business insights across banking and healthcare. Applies modern analytics and Generative AI-assisted workflows to accelerate development while independently validating final logic and outputs.")
p.paragraph_format.space_after=Pt(1); p.paragraph_format.line_spacing=1.04; p.runs[0].font.size=Pt(8.6)

heading(doc,"Technical Skills")
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(1); p.paragraph_format.line_spacing=1.0
items=[("Data Analytics","SQL, Python, Pandas, NumPy, Excel, data cleaning, transformation, validation, EDA, statistical analysis, KPI analysis, business analysis"),("Business Intelligence","Power BI, DAX, Power Query, data visualization, dashboard development, dimensional modeling, data storytelling"),("Database","PostgreSQL"),("Modern Analytics","PySpark, Apache Spark, Delta Lake, dbt, medallion architecture, Airflow DAG design, incremental processing, automated data quality"),("AI-Assisted Analytics","ChatGPT, Microsoft Copilot, SQL/Python and formula troubleshooting, workflow optimization, documentation"),("Tools","Git, GitHub, ArcGIS Pro, Google Earth Pro")]
for i,(a,b) in enumerate(items):
    rr=p.add_run(f"{a}: "); rr.bold=True; rr.font.size=Pt(8.35); rr=p.add_run(b+("  |  " if i<len(items)-1 else "")); rr.font.size=Pt(8.35)

heading(doc,"Professional Experience")
role(doc,"Senior Associate","Tech Mahindra","Mar 2026 - Present")
bullet(doc,"Analyze and validate high-volume structured operational and location-based datasets against defined specifications and business rules, identifying inconsistencies and maintaining accuracy.")
bullet(doc,"Execute quality controls and consistency checks across complex data attributes to detect exceptions, incomplete records and conflicting values.")
bullet(doc,"Investigate edge cases and root causes, document findings and coordinate correction workflows to maintain reliable analytical datasets.")
role(doc,"Analyst","Blue Dome Technologies","Jan 2024 - Dec 2025")
bullet(doc,"Extracted, cleaned, transformed and validated multi-source structured datasets using Python, Pandas, Excel and GIS tools, producing analysis-ready outputs.")
bullet(doc,"Profiled and reconciled asset attributes, GPS/KML records and source imagery to identify missing, duplicate, anomalous and inconsistent records.")
bullet(doc,"Applied business-rule validation and quality audits, prepared analytical datasets and structured reports, and coordinated resolution of recurring data exceptions.")

heading(doc,"Selected Projects")
role(doc,"FinTrust Banking Intelligence","End-to-End Banking Analytics & Modern BI Platform","2026")
bullet(doc,"Unified customer, transaction, deposit, lending, fraud, collection, churn and branch data to support management analysis of behavior, trends, portfolio risk and operating performance.")
bullet(doc,"Cleaned, joined, transformed and validated 257,969 records; used SQL and Python/Pandas to calculate and reconcile KPIs across 5,000 customers and 120,000 transactions.")
bullet(doc,"Built a six-page interactive Power BI solution using Power Query, dimensional modeling, DAX measures, synchronized filters and decision-focused trend and drill-down analysis.")
bullet(doc,"Implemented a local PySpark/Delta Lake Bronze-Silver-Gold pipeline with incremental processing, 11 dbt models, 12 passing dbt tests, 8 repository tests and an eight-task Airflow DAG design.")
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(1); p.add_run("Repository: ").bold=True; hyperlink(p,"github.com/Janakiram10/FinTrust-Banking-Intelligence","https://github.com/Janakiram10/FinTrust-Banking-Intelligence")
role(doc,"MediCore Healthcare Analytics","Healthcare Revenue, Patient Flow & Operational Intelligence","2026")
bullet(doc,"Integrated and analyzed 500,000 appointments and 300,000 encounters using PostgreSQL, SQL and Python/Pandas to create a centralized view of patient flow and hospital performance.")
bullet(doc,"Cleaned and validated healthcare data, performed EDA and calculated KPIs including 9.81% no-shows, 62.86% collections, 62.18-minute waits, 74.61 satisfaction and 3.65% readmissions.")
bullet(doc,"Built Power BI and Plotly Dash reporting for revenue and collections (₹178.13 Cr billed; ₹111.98 Cr collected), no-shows, patient experience, wait times and readmission risk.")
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(1); p.add_run("Repository: ").bold=True; hyperlink(p,"github.com/Janakiram10/MediCore-Healthcare-Analytics","https://github.com/Janakiram10/MediCore-Healthcare-Analytics")

heading(doc,"Education")
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0
r=p.add_run("Bachelor of Technology in Mechanical Engineering"); r.bold=True; r.font.size=Pt(8.6)
p.add_run(" | WISTM, Andhra University | Completed 2023 | GPA: 7.19/10").font.size=Pt(8.5)

for table in doc.tables:
    table.allow_autofit=False
doc.save(DOCX); print(DOCX)
